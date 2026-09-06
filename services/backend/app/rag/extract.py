"""Text extraction from uploaded files (Phase 5).

Deliberately small and format-limited: PDF, DOCX, TXT and Markdown cover
what a person actually uploads to a personal assistant, and each extractor
here is a pure function over bytes so the whole module is testable without a
filesystem or a running server.

Page numbers are the reason this is not just "decode the bytes". A citation
that says "page 4 of handbook.pdf" is one the reader can act on; a citation
that says "somewhere in handbook.pdf" is barely better than none
(docs/architecture.md Section 7). So extraction returns pages, not a string,
and formats that genuinely have no pages say so with `page_number=None`
rather than inventing one.
"""
from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass

#: Upload ceiling. Not a technical limit -- the real constraint is that
#: embedding a large document costs free-tier quota and takes long enough
#: that a user assumes it has hung. Raise it deliberately, not by accident.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".markdown")


class ExtractionError(Exception):
    """Raised when a file cannot be turned into text.

    Always carries a message meant for the user rather than a stack trace:
    it ends up in `documents.error` and on screen.
    """


@dataclass(frozen=True)
class Page:
    #: 1-based, or None for formats with no concept of a page.
    number: int | None
    text: str


def content_hash(pages: list[Page]) -> str:
    """sha256 of the extracted text, whitespace-normalised.

    Hashing the *text* rather than the uploaded bytes on purpose: the same
    document re-saved by a different PDF writer has different bytes and
    identical content, and storing it twice would put two copies of every
    passage into retrieval, competing with each other for the same slots.
    """
    normalised = " ".join(" ".join(p.text.split()) for p in pages)
    return hashlib.sha256(normalised.encode("utf-8")).hexdigest()


def _extract_pdf(data: bytes) -> list[Page]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ExtractionError("PDF support needs the `pypdf` package to be installed.") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"That PDF could not be opened ({type(exc).__name__}).") from exc

    if getattr(reader, "is_encrypted", False):
        # Trying and failing later is worse than saying so now.
        raise ExtractionError("That PDF is password-protected. Remove the password and upload it again.")

    pages: list[Page] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            # One unreadable page should not lose the other ninety-nine.
            text = ""
        pages.append(Page(number=index, text=text))

    if not any(p.text.strip() for p in pages):
        # The single most common surprise with PDFs, and worth naming
        # precisely: a scan is an image, and there is no text layer to find.
        raise ExtractionError(
            "No text could be read from that PDF. If it is a scan, it needs OCR first -- "
            "CIPHER does not do OCR yet."
        )
    return pages


def _extract_docx(data: bytes) -> list[Page]:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ExtractionError("DOCX support needs the `python-docx` package to be installed.") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"That DOCX could not be opened ({type(exc).__name__}).") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise ExtractionError("That DOCX appears to contain no text.")
    # DOCX has no reliable page boundaries without rendering it, so this
    # honestly reports one pageless unit rather than guessing at page breaks.
    return [Page(number=None, text=text)]


def _extract_text(data: bytes) -> list[Page]:
    for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes any byte string
        raise ExtractionError("That file is not readable as text.")

    if not text.strip():
        raise ExtractionError("That file is empty.")
    return [Page(number=None, text=text)]


def extract(filename: str, data: bytes) -> list[Page]:
    """Turn an uploaded file into pages of text, or raise ExtractionError."""
    if not data:
        raise ExtractionError("That file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError(
            f"That file is {len(data) / 1_048_576:.1f} MB; the limit is "
            f"{MAX_UPLOAD_BYTES // 1_048_576} MB."
        )

    lowered = filename.lower()
    if lowered.endswith(".pdf"):
        return _extract_pdf(data)
    if lowered.endswith(".docx"):
        return _extract_docx(data)
    if lowered.endswith((".txt", ".md", ".markdown")):
        return _extract_text(data)

    if lowered.endswith(".doc"):
        # Named separately because "why did .docx work and .doc not" is an
        # obvious question with a real answer.
        raise ExtractionError("Legacy .doc files are not supported. Save it as .docx and try again.")
    raise ExtractionError(
        f"Unsupported file type. Supported: {', '.join(SUPPORTED_EXTENSIONS)}."
    )
